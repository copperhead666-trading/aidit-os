#!/usr/bin/env python3
"""Generate Template Surat Jalan .xlsx for katering ops."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

wb = Workbook()
ws = wb.active
ws.title = "Surat Jalan"

bold = Font(bold=True, size=11)
bold_big = Font(bold=True, size=16)
bold_med = Font(bold=True, size=13)
align_right = Alignment(horizontal='right')
align_center = Alignment(horizontal='center')
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
fill_header = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
font_white = Font(bold=True, color="FFFFFF", size=11)

ws.column_dimensions['A'].width = 5
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 14
ws.column_dimensions['E'].width = 22

# === HEADER ===
ws.merge_cells('B2:E2')
ws['B2'] = 'KATERING SEJAHTERA'
ws['B2'].font = bold_big
ws['B2'].alignment = Alignment(horizontal='center')

ws.merge_cells('B3:E3')
ws['B3'] = 'Jl. Contoh No. 123, Jakarta | Telp: 021-12345678'
ws['B3'].font = Font(size=9, italic=True)
ws['B3'].alignment = Alignment(horizontal='center')

ws.merge_cells('B5:E5')
ws['B5'] = 'SURAT JALAN'
ws['B5'].font = bold_big
ws['B5'].alignment = Alignment(horizontal='center')

# Meta
ws['B7'] = 'No. Surat Jalan:'
ws['B7'].font = bold
ws.merge_cells('C7:E7')
ws['C7'] = '=CONCATENATE("SJ-",TEXT(D8,"YYYYMMDD"),"-",TEXT(1,"0000"))'
ws['C7'].font = Font(color="0000FF", underline="single")

ws['B8'] = 'Tanggal Kirim:'
ws['B8'].font = bold
ws['D8'] = ''
ws['D8'].number_format = 'DD/MM/YYYY'

ws['B9'] = 'Tujuan Pengiriman:'
ws['B9'].font = bold
ws.merge_cells('C9:E9')

ws['B10'] = 'No. Telepon Tujuan:'
ws['B10'].font = bold
ws.merge_cells('C10:E10')

# === ITEM TABLE ===
row = 12
ws.merge_cells(f'B{row}:E{row}')
ws[f'B{row}'] = 'Daftar Barang Dikirim'
ws[f'B{row}'].font = bold_med
ws[f'B{row}'].alignment = Alignment(horizontal='center')

row = 13
headers = ['No', 'Nama Item', 'Jumlah', 'Satuan', 'Catatan']
cols = ['A', 'B', 'C', 'D', 'E']
for c, hdr in zip(cols, headers):
    cell = ws[f'{c}{row}']
    cell.value = hdr
    cell.font = font_white
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = thin_border

for idx in range(1, 16):
    r = 13 + idx
    ws[f'A{r}'] = idx
    ws[f'A{r}'].alignment = align_center
    ws[f'A{r}'].border = thin_border
    ws[f'B{r}'].border = thin_border
    ws[f'C{r}'].border = thin_border
    ws[f'C{r}'].alignment = align_center
    ws[f'D{r}'].border = thin_border
    ws[f'D{r}'].alignment = align_center
    ws[f'E{r}'].border = thin_border

# === DRIVER INFO ===
r = 30
ws.merge_cells(f'B{r}:E{r}')
ws[f'B{r}'] = 'Informasi Pengiriman'
ws[f'B{r}'].font = bold_med

ws[f'B{r+1}'] = 'Nama Pengemudi:'
ws[f'B{r+1}'].font = bold
ws.merge_cells(f'C{r+1}:E{r+1}')

ws[f'B{r+2}'] = 'No. HP Pengemudi:'
ws[f'B{r+2}'].font = bold
ws.merge_cells(f'C{r+2}:E{r+2}')

ws[f'B{r+3}'] = 'Kendaraan:'
ws[f'B{r+3}'].font = bold
ws.merge_cells(f'C{r+3}:E{r+3}')

ws[f'B{r+4}'] = 'No. Polisi:'
ws[f'B{r+4}'].font = bold
ws.merge_cells(f'C{r+4}:E{r+4}')

# === SIGNATURES ===
r2 = r + 6
ws[f'B{r2}'] = 'Pengirim:'
ws[f'B{r2}'].font = bold
ws[f'D{r2}'] = 'Penerima:'
ws[f'D{r2}'].font = bold

ws[f'B{r2+1}'] = '(                                        )'
ws[f'D{r2+1}'] = '(                                        )'

ws[f'B{r2+3}'] = 'Tanggal: ____/____/________'
ws[f'D{r2+3}'] = 'Tanggal: ____/____/________'

# Stamp area note
ws.merge_cells(f'B{r2+5}:E{r2+5}')
ws[f'B{r2+5}'] = '* Stempel perusahaan di area tanda tangan pengirim'
ws[f'B{r2+5}'].font = Font(size=9, italic=True)

ws.print_area = 'A1:E40'

wb.save('D:/AI/aidit-os-v5/docs/deliverables/katering-ops/template_surat_jalan.xlsx')
print("template_surat_jalan.xlsx saved")